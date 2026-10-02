import io
from datetime import date

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from sqlalchemy.orm import Session, selectinload

from app.core.i18n import t
from app.core.tz import user_tz
from app.models.bill import BillFrequency, BillTemplate, PaymentInstance
from app.models.user import User
from app.schemas.auth import EXPORT_FIELD_KEYS
from app.services.notify import footer_text
from app.services.recurrence import backfill_template_instances

_COLUMNS = [
    "Bill",
    "Category",
    "Period",
    "Due Date",
    "Amount",
    "Currency",
    "Status",
    "Paid Amount",
    "Paid At",
    "Notes",
]

# Index-aligned with _COLUMNS: EXPORT_FIELD_KEYS[i] is the stable identifier for
# _COLUMNS[i], stored on User.export_fields (snake_case, locale-independent).
_COLUMN_KEYS = list(EXPORT_FIELD_KEYS)

# Mirrors the status colors used in the frontend (PaymentsCalendar.tsx STATUS_TILE).
_STATUS_STYLES: dict[str, tuple[str, str]] = {
    "upcoming": ("DBEAFE", "1D4ED8"),  # blue-100 / blue-700
    "overdue": ("FEE2E2", "B91C1C"),  # red-100 / red-700
    "paid": ("DCFCE7", "15803D"),  # green-100 / green-700
}
_HEADER_FILL = PatternFill("solid", fgColor="F3F4F6")  # gray-100
_HEADER_FONT = Font(bold=True)
_MAX_COL_WIDTH = 40

# Amount/Paid Amount are written as Text cells (not Number) using the user's chosen
# decimal_separator, so the value renders identically regardless of the opening
# machine's regional Excel locale. A "[$-LCID]0.00" Number format was tried instead
# (to avoid Excel's "stored as text" warning) but the locale tag does NOT actually
# control which glyph a plain numeric placeholder renders with — that's driven by
# the opening machine's own regional settings regardless of any locale tag — so a
# Polish-locale Excel showed a comma no matter which format was applied. Confirmed
# working as Text; do not switch back without verifying in a real, non-US-locale
# Excel first.
_AMOUNT_TEXT_FORMAT = "@"

# Mirrors CATEGORY_COLOR_BORDER's light-mode shades in frontend/src/lib/categories.ts —
# a left-border accent instead of a full cell fill, so many categories don't turn
# the sheet into a rainbow. Falls back to slate, same as the frontend's FALLBACK_COLOR.
_CATEGORY_BORDER_HEX: dict[str, str] = {
    "blue": "60A5FA",
    "purple": "C084FC",
    "rose": "FB7185",
    "orange": "FB923C",
    "slate": "94A3B8",
    "violet": "A78BFA",
    "cyan": "06B6D4",
    "emerald": "34D399",
    "slate-light": "CBD5E1",
    "yellow": "FACC15",
    "lime": "A3E635",
    "pink": "F472B6",
    "blue-dark": "1D4ED8",
    "emerald-dark": "047857",
    "rose-dark": "BE123C",
    "orange-dark": "9A3412",
}
_CATEGORY_BORDER_FALLBACK = _CATEGORY_BORDER_HEX["slate"]


def _ensure_year_instances(db: Session, user_id: int, year: int) -> None:
    """Backfill missing payment instances for every eligible template across
    the full year, so export isn't limited to months the user has already
    visited in the UI (list_payments/sync-instances only seed on demand)."""
    templates = (
        db.query(BillTemplate)
        .filter(
            BillTemplate.user_id == user_id,
            BillTemplate.is_archived.is_(False),
            BillTemplate.is_paused.is_(False),
            BillTemplate.frequency != BillFrequency.one_off,
        )
        .all()
    )
    for template in templates:
        backfill_template_instances(db, template, f"{year}-01", f"{year}-12")


def _rows_by_month(
    db: Session, me: User, lang: str, year: int
) -> dict[int, list[dict]]:
    _ensure_year_instances(db, me.id, year)
    tz = user_tz(me)

    instances = (
        db.query(PaymentInstance)
        .options(selectinload(PaymentInstance.template))
        .join(BillTemplate, PaymentInstance.bill_id == BillTemplate.id)
        .filter(
            BillTemplate.user_id == me.id,
            PaymentInstance.period.startswith(f"{year}-"),
            PaymentInstance.is_deleted.is_(False),
        )
        .order_by(PaymentInstance.due_date)
        .all()
    )

    # Index instances by month number (1–12)
    by_month: dict[int, list[dict]] = {m: [] for m in range(1, 13)}
    for i in instances:
        by_month[int(i.period[5:7])].append(
            {
                "Bill": i.template.name,
                "Category": i.template.category.name,
                "Period": i.period,
                "Due Date": i.due_date.isoformat(),
                "Amount": f"{i.current_amount:.2f}".replace(".", me.decimal_separator),
                "Currency": i.template.currency,
                "Status": t(lang, f"PaymentRow.status.{i.status}"),
                "Paid Amount": (
                    f"{i.paid_amount:.2f}".replace(".", me.decimal_separator)
                    if i.paid_amount
                    else None
                ),
                "Paid At": (
                    i.paid_at.astimezone(tz).date().isoformat() if i.paid_at else None
                ),
                "Notes": i.notes,
                "_status_key": i.status,
                "_category_color": i.template.category.color,
            }
        )
    return by_month


def build_xlsx(
    db: Session, me: User, today: date, year: int, month: int | None, lang: str
) -> io.BytesIO:
    selected_keys = set(me.export_fields)
    selected_columns = [
        col for key, col in zip(_COLUMN_KEYS, _COLUMNS) if key in selected_keys
    ]

    by_month = _rows_by_month(db, me, lang, year)

    headers = [
        t(lang, f"SettingsPage.excelExport.fields.{key}")
        for key in _COLUMN_KEYS
        if key in selected_keys
    ]
    status_col_idx = (
        selected_columns.index("Status") + 1 if "Status" in selected_columns else None
    )
    category_col_idx = (
        selected_columns.index("Category") + 1
        if "Category" in selected_columns
        else None
    )
    amount_col_idxs = {
        selected_columns.index(c) + 1
        for c in ("Amount", "Paid Amount")
        if c in selected_columns
    }

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for m in [month] if month else range(1, 13):
            sheet_name = f"{t(lang, f"ExcelExport.monthShort.{m}")} {year}"
            rows = by_month[m]
            df = (
                pd.DataFrame(rows, columns=selected_columns)
                if rows
                else pd.DataFrame(columns=selected_columns)
            )
            df.columns = headers
            df.to_excel(writer, index=False, sheet_name=sheet_name)

            ws = writer.sheets[sheet_name]
            ws.page_setup.orientation = "landscape"
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0  # as many pages tall as needed
            ws.oddFooter.left.text = f"{footer_text()} · {today.isoformat()}"
            ws.oddFooter.right.text = "&P / &N"
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
                max_len = len(header)
                for row_idx in range(2, len(rows) + 2):
                    value = ws.cell(row=row_idx, column=col_idx).value
                    if value is not None:
                        max_len = max(max_len, len(str(value)))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(
                    max_len + 2, _MAX_COL_WIDTH
                )

            for row_idx, row in enumerate(rows, start=2):
                if status_col_idx is not None:
                    fill_color, font_color = _STATUS_STYLES.get(
                        row["_status_key"], (None, None)
                    )
                    if fill_color:
                        cell = ws.cell(row=row_idx, column=status_col_idx)
                        cell.fill = PatternFill("solid", fgColor=fill_color)
                        cell.font = Font(color=font_color)

                if category_col_idx is not None:
                    border_color = _CATEGORY_BORDER_HEX.get(
                        row["_category_color"], _CATEGORY_BORDER_FALLBACK
                    )
                    cell = ws.cell(row=row_idx, column=category_col_idx)
                    cell.border = Border(left=Side(style="thick", color=border_color))

                for amount_col_idx in amount_col_idxs:
                    ws.cell(row=row_idx, column=amount_col_idx).number_format = (
                        _AMOUNT_TEXT_FORMAT
                    )

        active_month = today.month if year == today.year else 1
        writer.book.active = 0 if month else active_month - 1
    buf.seek(0)
    return buf
