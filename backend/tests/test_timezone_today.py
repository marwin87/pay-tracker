"""'Today', the month and overdue follow the user's own time zone, not the server's."""

import io
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from tests.conftest import auth, category_id, register_and_login

# 23:30 UTC on 10 March 2026 is already the 11th in Auckland (UTC+13, daylight time)
# and still the 10th in Los Angeles (UTC-7) and UTC.
_CLOCK = datetime(2026, 3, 10, 23, 30, tzinfo=timezone.utc)


@pytest.fixture()
def pin(monkeypatch):
    def _pin(moment: datetime) -> None:
        monkeypatch.setattr("app.core.tz._utcnow", lambda: moment)

    _pin(_CLOCK)
    return _pin


def _user(client, email: str, zone: str) -> str:
    token = register_and_login(client, email)
    r = client.patch("/auth/me", json={"timezone": zone}, headers=auth(token))
    assert r.status_code == 200, r.text
    return token


def _bill(client, token, **overrides) -> dict:
    body = {
        "name": "Rent",
        "category_id": category_id(client, token),
        "frequency": "monthly",
        "amount": "100.00",
        "currency": "PLN",
        "due_day": 10,
        **overrides,
    }
    r = client.post("/bills", json=body, headers=auth(token))
    assert r.status_code == 201, r.text
    return r.json()


def _payments(client, token, month=None) -> list[dict]:
    qs = f"?month={month}" if month else ""
    r = client.get(f"/bills/payments{qs}", headers=auth(token))
    assert r.status_code == 200, r.text
    return r.json()


def _seed_march(client, token, **bill) -> dict:
    _bill(client, token, **bill)
    assert (
        client.post(
            "/bills/sync-instances?month=2026-03", headers=auth(token)
        ).status_code
        == 204
    )
    [payment] = _payments(client, token, "2026-03")
    return payment


def test_overdue_is_relative_to_the_users_today(client, pin):
    utc = _user(client, "utc@test.com", "UTC")
    nz = _user(client, "nz@test.com", "Pacific/Auckland")
    la = _user(client, "la@test.com", "America/Los_Angeles")

    # due on the 10th: still today for UTC and Los Angeles, yesterday for Auckland
    assert _seed_march(client, utc)["status"] == "upcoming"
    assert _seed_march(client, nz)["status"] == "overdue"
    assert _seed_march(client, la)["status"] == "upcoming"


def test_default_month_follows_the_users_zone(client, pin):
    pin(datetime(2026, 3, 31, 23, 30, tzinfo=timezone.utc))
    utc = _user(client, "m-utc@test.com", "UTC")
    nz = _user(client, "m-nz@test.com", "Pacific/Auckland")
    for token in (utc, nz):
        _bill(client, token)

    # sync + list with no month: March for UTC, already April in Auckland
    for token, month in ((utc, "2026-03"), (nz, "2026-04")):
        assert (
            client.post("/bills/sync-instances", headers=auth(token)).status_code == 204
        )
        assert [p["period"] for p in _payments(client, token)] == [month]


def test_new_bill_starts_in_the_users_current_month(client, pin):
    pin(datetime(2026, 3, 31, 23, 30, tzinfo=timezone.utc))
    utc = _user(client, "s-utc@test.com", "UTC")
    nz = _user(client, "s-nz@test.com", "Pacific/Auckland")
    assert _bill(client, utc)["start_period"] == "2026-03"
    assert _bill(client, nz)["start_period"] == "2026-04"


def test_future_payment_date_limit_is_tomorrow_in_the_users_zone(client, pin):
    # UTC today = 10 Mar (tomorrow 11 ok, 12 not); Auckland today = 11 Mar (12 ok, 13 not)
    for zone, ok, too_far in (
        ("UTC", "2026-03-11", "2026-03-12"),
        ("Pacific/Auckland", "2026-03-12", "2026-03-13"),
    ):
        token = _user(client, f"f-{zone[:3]}@test.com", zone)
        pid = _seed_march(client, token)["id"]
        pay = f"/bills/payments/{pid}/pay"
        assert (
            client.post(pay, json={"paid_at": too_far}, headers=auth(token)).status_code
            == 400
        )
        assert (
            client.post(pay, json={"paid_at": ok}, headers=auth(token)).status_code
            == 200
        )


def test_picked_payment_date_is_noon_in_the_users_zone(client, pin):
    token = _user(client, "noon@test.com", "Pacific/Auckland")
    pid = _seed_march(client, token)["id"]
    r = client.post(
        f"/bills/payments/{pid}/pay",
        json={"paid_at": "2026-03-05"},
        headers=auth(token),
    )
    assert r.status_code == 200
    paid = datetime.fromisoformat(r.json()["paid_at"])
    # noon on 5 March in Auckland (UTC+13) is 23:00 UTC on the 4th
    assert paid == datetime(2026, 3, 4, 23, 0, tzinfo=timezone.utc)
    local = paid.astimezone(ZoneInfo("Pacific/Auckland"))
    assert (local.date().isoformat(), local.hour) == ("2026-03-05", 12)


def _sheet_values(response) -> set:
    wb = load_workbook(io.BytesIO(response.content))
    return {c.value for row in wb.worksheets[0].iter_rows() for c in row}


def test_export_shows_the_paid_date_in_the_users_zone(client, pin):
    # paid "now" = 23:30 UTC on the 10th = 11 March in Auckland
    token = _user(client, "x@test.com", "Pacific/Auckland")
    pid = _seed_march(client, token, due_day=20)["id"]
    client.post(f"/bills/payments/{pid}/pay", json={}, headers=auth(token))
    r = client.get("/export/xlsx?year=2026&month=3", headers=auth(token))
    assert r.status_code == 200
    values = _sheet_values(r)
    assert "2026-03-11" in values and "2026-03-10" not in values


def test_export_default_year_and_backup_name_follow_the_users_zone(client, pin):
    pin(datetime(2025, 12, 31, 23, 30, tzinfo=timezone.utc))
    utc = _user(client, "y-utc@test.com", "UTC")
    nz = _user(client, "y-nz@test.com", "Pacific/Auckland")

    def disposition(token, path):
        return client.get(path, headers=auth(token)).headers["content-disposition"]

    assert "2025.xlsx" in disposition(utc, "/export/xlsx")
    assert "2026.xlsx" in disposition(nz, "/export/xlsx")
    assert "backup-2025-12-31.bak" in disposition(utc, "/export/json")
    assert "backup-2026-01-01.bak" in disposition(nz, "/export/json")


def test_unpay_status_uses_the_users_today(client, pin):
    token = _user(client, "unpay@test.com", "Pacific/Auckland")
    pid = _seed_march(client, token)["id"]  # due the 10th: overdue in Auckland
    client.post(f"/bills/payments/{pid}/pay", json={}, headers=auth(token))
    r = client.post(f"/bills/payments/{pid}/unpay", headers=auth(token))
    assert r.json()["status"] == "overdue"
