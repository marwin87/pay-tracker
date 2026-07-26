import io
from decimal import Decimal

import openpyxl

from paytracker.data.models import BillCategory, BillFrequency
from paytracker.services import bills, payments, xlsx_export


def test_export_xlsx_produces_valid_workbook_with_12_sheets(db_session):
    bill = bills.create_bill(
        db_session,
        name="Rent",
        category=BillCategory.housing,
        frequency=BillFrequency.monthly,
        amount=Decimal("1000.00"),
        currency="EUR",
        due_day=1,
    )
    payments.sync_instances(db_session)

    data = xlsx_export.export_xlsx_bytes(db_session, year=payments.list_payments(db_session)[0].due_date.year)

    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) == 12

    # exactly one sheet has the seeded row
    row_counts = [ws.max_row for ws in wb.worksheets]
    assert max(row_counts) == 2  # header + 1 data row
    assert sum(1 for c in row_counts if c == 2) == 1
